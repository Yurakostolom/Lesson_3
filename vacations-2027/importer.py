"""Read employee rosters and match explicitly selected people without guessing names.

Uploaded documents are data only. No attachment is executed or retained here.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_ROWS = 25_000
MAX_UNPACKED_SIZE = 100 * 1024 * 1024
SUPPORTED_SUFFIXES = frozenset({".xlsx", ".csv", ".txt", ".docx", ".pdf"})

_LETTERS = "А-Яа-яЁёA-Za-z"
_NAME = rf"[А-ЯЁA-Z][{_LETTERS}]+(?:[-'’][{_LETTERS}]+)*"
_FULL_NAME = re.compile(rf"^({_NAME})\s+({_NAME})\s+({_NAME})(?=$|[\s#\d,;|])")
_INITIALS = re.compile(rf"^({_NAME})\s+([А-ЯЁA-Z])\.?(?:\s*([А-ЯЁA-Z])\.?)?$")
_BLOCKED_STARTS = {
    "врач", "врачи", "дата", "дневной", "должность", "итого", "всего", "имя",
    "кабинет", "категория", "контакты", "логин", "манипуляционная", "месяц",
    "отдел", "отпуск", "отпуска", "отчество", "план", "подразделение",
    "процедурный", "рентген", "сотрудник", "сотрудники", "список", "стажер",
    "тестовый", "фамилия", "фио", "физиотерапия", "функциональная", "январь",
    "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь",
    "октябрь", "ноябрь", "декабрь", "ответственный", "примечание", "примечания",
    "surname", "employee", "name", "total", "first", "last", "full",
}
_SURNAME_HEADERS = {"фамилия", "surname", "last name", "lastname"}
_GIVEN_HEADERS = {"имя", "first name", "firstname", "given name"}
_MIDDLE_HEADERS = {"отчество", "middle name", "patronymic"}
_PERSON_HEADERS = {"фио", "ф и о", "фамилия имя отчество", "сотрудник", "сотрудники", "врач", "врачи", "employee", "full name", "name"}


def _key(value: str) -> str:
    return unicodedata.normalize("NFC", value).casefold().replace("ё", "е")


def _text(value: object) -> str:
    text = unicodedata.normalize("NFC", str(value or ""))
    text = text.replace("\u00a0", " ").replace("\u200b", "").replace("\ufeff", "")
    # The PDF export can omit spaces between the surname, name and patronymic.
    # This restores explicit capital-letter boundaries; it never corrects spelling.
    text = re.sub(r"(?<=[а-яёa-z])(?=[А-ЯЁA-Z])", " ", text)
    return " ".join(text.split()).strip()


def _without_number(value: str) -> str:
    return re.sub(r"^\s*\d+[.)]?\s+", "", _text(value))


def _header(value: str) -> str:
    return re.sub(r"[.():№]+", " ", _key(_text(value))).strip()


def _blocked(value: str) -> bool:
    first = value.split(maxsplit=1)[0] if value else ""
    return _key(first) in _BLOCKED_STARTS


def _check_archive(path: Path) -> None:
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10_000 or sum(entry.file_size for entry in entries) > MAX_UNPACKED_SIZE:
            raise ValueError("Файл содержит слишком большой объём распакованных данных (максимум 100 МБ).")


def _read_rows(path: Path) -> list[list[str]]:
    path = Path(path)
    if not path.is_file():
        raise ValueError("Файл для импорта не найден. Загрузите документ заново.")
    if path.stat().st_size == 0:
        raise ValueError("Загружен пустой файл.")
    if path.stat().st_size > MAX_FILE_SIZE:
        raise ValueError("Размер каждого файла не должен превышать 10 МБ.")
    suffix = path.suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("Поддерживаются только XLSX, CSV, TXT, DOCX и PDF.")
    rows: list[list[str]] = []
    try:
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            _check_archive(path)
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in workbook:
                    for values in sheet.iter_rows(values_only=True):
                        rows.append([_text(value) for value in values])
                        if len(rows) > MAX_ROWS:
                            raise ValueError("Слишком много строк в документе (максимум 25 000).")
                    rows.append([])
            finally:
                workbook.close()
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("PDF защищён паролем. Загрузите документ без защиты.")
            for page in reader.pages:
                rows.extend([[line] for line in (page.extract_text() or "").splitlines()])
            if not any(any(cell.strip() for cell in row) for row in rows):
                raise ValueError("В PDF нет распознаваемого текста. Загрузите текстовый PDF или XLSX, CSV, TXT, DOCX.")
        elif suffix == ".docx":
            _check_archive(path)
            with ZipFile(path) as archive:
                tree = ElementTree.fromstring(archive.read("word/document.xml"))
            namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            body = tree.find("w:body", namespace)
            if body is not None:
                for node in body:
                    if node.tag.endswith("}tbl"):
                        for row in node.findall("w:tr", namespace):
                            rows.append([" ".join(element.text or "" for element in cell.findall(".//w:t", namespace)) for cell in row.findall("w:tc", namespace)])
                    elif node.tag.endswith("}p"):
                        rows.append(["".join(element.text or "" for element in node.findall(".//w:t", namespace))])
        else:
            raw = path.read_bytes()
            try:
                decoded = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                decoded = raw.decode("cp1251")
            if suffix == ".csv":
                try:
                    dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t|")
                except csv.Error:
                    rows = list(csv.reader(io.StringIO(decoded), delimiter=";"))
                else:
                    rows = list(csv.reader(io.StringIO(decoded), dialect))
            else:
                rows = [re.split(r"\t|;", line) for line in decoded.splitlines()]
    except ValueError:
        raise
    except (OSError, BadZipFile, KeyError, ElementTree.ParseError, UnicodeError, csv.Error) as exc:
        raise ValueError("Не удалось прочитать документ. Проверьте формат и загрузите файл заново.") from exc
    except Exception as exc:
        # Parser errors can contain source contents; keep the public message generic.
        raise ValueError("Не удалось прочитать документ. Проверьте, что файл не повреждён и не защищён паролем.") from exc
    if len(rows) > MAX_ROWS:
        raise ValueError("Слишком много строк в документе (максимум 25 000).")
    return [[_text(cell) for cell in row] for row in rows]


def _person_columns(row: list[str]) -> tuple[int, ...] | None:
    headers = [_header(cell) for cell in row]
    surname = next((i for i, value in enumerate(headers) if value in _SURNAME_HEADERS), None)
    given = next((i for i, value in enumerate(headers) if value in _GIVEN_HEADERS), None)
    middle = next((i for i, value in enumerate(headers) if value in _MIDDLE_HEADERS), None)
    if surname is not None:
        return tuple(i for i in (surname, given, middle) if i is not None)
    person = next((i for i, value in enumerate(headers) if value in _PERSON_HEADERS), None)
    return (person,) if person is not None else None


def _row_person_values(rows: list[list[str]]) -> list[str]:
    """Respect labelled columns, otherwise read the first plausible person cell."""
    values: list[str] = []
    columns: tuple[int, ...] | None = None
    for row in rows:
        if not row:
            # Headers may change between workbook sheets.
            columns = None
            continue
        if not any(row):
            continue
        detected = _person_columns(row)
        strong_header = any(_header(cell) in (_SURNAME_HEADERS | _GIVEN_HEADERS | _MIDDLE_HEADERS | {"фио", "ф и о", "full name"}) for cell in row)
        if detected is not None and (columns is None or strong_header):
            columns = detected
            continue
        if columns is not None:
            values.append(" ".join(row[i] for i in columns if i < len(row)))
            continue
        nonempty = [cell for cell in row if cell]
        if not nonempty:
            continue
        # A leading numeric table index is not part of the employee name.
        if re.fullmatch(r"\d+[.)]?", nonempty[0]):
            nonempty = nonempty[1:]
        if nonempty:
            values.append(" ".join(nonempty[:3]))
    return values


def _full_name(value: str) -> str | None:
    value = _without_number(value)
    if _blocked(value):
        return None
    match = _FULL_NAME.match(value)
    if match:
        return " ".join(match.groups())
    # Plain text may use all lowercase names; require a complete three-word FIO.
    if re.fullmatch(rf"{_NAME}\s+{_NAME}\s+{_NAME}", value, re.IGNORECASE):
        return value
    return None


@dataclass(frozen=True)
class _Selection:
    surname: str
    parts: tuple[str, ...]
    display: str


def _selection(value: str) -> _Selection | None:
    value = _without_number(value)
    if not value or _blocked(value):
        return None
    initial_match = _INITIALS.fullmatch(value)
    if initial_match:
        parts = tuple(part for part in initial_match.groups() if part)
        return _Selection(parts[0], parts, value)
    parts = tuple(value.split())
    if 1 <= len(parts) <= 3 and all(re.fullmatch(_NAME, part, re.IGNORECASE) for part in parts):
        return _Selection(parts[0], parts, " ".join(parts))
    # For PDF table rows, ignore the metadata following a complete FIO.
    full = _full_name(value)
    if full:
        parts = tuple(full.split())
        return _Selection(parts[0], parts, full)
    return None


def _matches(selection: _Selection, full_name: str) -> bool:
    candidate_parts = full_name.split()
    if len(candidate_parts) < len(selection.parts):
        return False
    for selected, candidate in zip(selection.parts, candidate_parts):
        if len(selected) == 1:
            if _key(selected) != _key(candidate[:1]):
                return False
        elif _key(selected) != _key(candidate):
            return False
    return True


def preview_import(vacations_path: Path, names_path: Path) -> dict:
    """Preview a join by surname and explicit initials/full names, never fuzzy match.

    Only entries in the first document can appear in ``matched``. A shared
    surname stays ambiguous unless the selection supplies identifying names or
    initials. The second document is authoritative for the resulting full FIO.
    """
    selected: list[_Selection] = []
    selection_keys: set[tuple[str, ...]] = set()
    warnings: list[str] = []
    for value in _row_person_values(_read_rows(Path(vacations_path))):
        person = _selection(value)
        if person is None:
            continue
        identity = tuple(_key(part) for part in person.parts)
        if identity not in selection_keys:
            selected.append(person)
            selection_keys.add(identity)
    if not selected:
        raise ValueError("В файле отпусков не найдены фамилии. Используйте столбец «Фамилия», «ФИО», «Врач» или список фамилий по одной на строку.")

    by_surname: dict[str, dict[str, str]] = {}
    for value in _row_person_values(_read_rows(Path(names_path))):
        full_name = _full_name(value)
        if full_name:
            surname_key = _key(full_name.split()[0])
            by_surname.setdefault(surname_key, {}).setdefault(_key(full_name), full_name)
    if not by_surname:
        raise ValueError("Во втором документе не найдены полные ФИО. Нужны фамилия, имя и отчество в одной строке или отдельных столбцах.")

    matched: list[dict[str, str]] = []
    unmatched: list[str] = []
    ambiguous: list[dict] = []
    matched_keys: set[str] = set()
    duplicate_matches = 0
    for person in selected:
        candidates = list(by_surname.get(_key(person.surname), {}).values())
        candidates = [candidate for candidate in candidates if _matches(person, candidate)]
        if not candidates:
            unmatched.append(person.display)
        elif len(candidates) > 1:
            ambiguous.append({"surname": person.surname, "candidates": sorted(candidates, key=_key)})
        else:
            full_name = candidates[0]
            identity = _key(full_name)
            if identity not in matched_keys:
                matched.append({"full_name": full_name, "surname": full_name.split()[0]})
                matched_keys.add(identity)
            else:
                duplicate_matches += 1
    if duplicate_matches:
        warnings.append(f"Повторные записи одного сотрудника объединены: {duplicate_matches}.")
    return {
        "matched": matched,
        "unmatched": unmatched,
        "ambiguous": ambiguous,
        "warnings": warnings,
        "counts": {
            "selected": len(matched) + len(unmatched) + len(ambiguous),
            "matched": len(matched),
            "unmatched": len(unmatched),
            "ambiguous": len(ambiguous),
        },
    }
